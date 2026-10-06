"""C2 document contracts first; real Docling, not a converter-shaped double."""

import asyncio
import hashlib
import io
import sys
from pathlib import Path

import pytest
from pydantic import ValidationError

from chimera.config import ChimeraConfig
from chimera.document_config import DocumentExtractionConfig
from chimera.documents import DocumentExtractor
from chimera.models import Extracted, Page
from chimera.refusals import ChimeraRefused, RefusalCode


def config(tmp_path, **updates):
    data = dict(
        schema="chimera.document-extraction/1",
        worker_python=sys.executable,
        work_directory=str(tmp_path / "worker"),
        max_input_bytes=2097152,
        max_output_bytes=2097152,
        max_diagnostic_bytes=65536,
        max_text_chars=200000,
        max_links=100,
        max_workers=2,
        timeout_seconds=30.0,
        cleanup_timeout_seconds=5.0,
        max_pages=100,
        max_archive_entries=1000,
        max_expanded_bytes=8388608,
        max_compression_ratio=1000.0,
        cpu_threads=2,
        pdf_pipeline="native",
        artifacts_directory=None,
        artifacts=(),
        do_ocr=False,
        do_table_structure=True,
        languages=("en", "zh", "ja", "ko", "id", "ms", "ru"),
        language_max_chars=8000,
        language_min_chars=30,
        language_min_confidence=0.5,
        language_min_margin=0.1,
    )
    data.update(updates)
    policy = DocumentExtractionConfig.model_validate(data)
    raw = ChimeraConfig.from_toml(Path("examples/chimera.toml")).model_dump(by_alias=True)
    raw["document_extraction"] = policy.model_dump(by_alias=True)
    return ChimeraConfig.model_validate(raw)


def docx(text=None):
    from docx import Document

    source = Document()
    source.add_heading("Port infrastructure report", level=1)
    source.add_paragraph(
        text
        or (
            "The port authority published a detailed report about maritime transport and "
            "infrastructure investment. Analysts reviewed documentary evidence describing "
            "terminal construction and shipping capacity."
        )
    )
    table = source.add_table(rows=2, cols=2)
    for cell, value in zip(table.rows[0].cells, ("Location", "Capacity"), strict=True):
        cell.text = value
    for cell, value in zip(table.rows[1].cells, ("Taiwan", "42"), strict=True):
        cell.text = value
    source.add_paragraph("Original report: https://example.org/appendix.pdf")
    stream = io.BytesIO()
    source.save(stream)
    return stream.getvalue()


def page(
    body=None,
    content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
):
    return Page(
        url="https://example.org/report.docx",
        final_url="https://example.org/report.docx",
        status=200,
        content_type=content_type,
        body=docx() if body is None else body,
    )


def test_real_docling_docx_retains_native_text_tables_structure_and_source_binding(tmp_path):
    source = page()
    cfg = config(tmp_path)
    result = asyncio.run(DocumentExtractor(cfg).extract(source))
    assert "Port infrastructure report" in result.title
    assert "Taiwan" in result.text and "42" in result.text
    assert result.language == "en"
    assert result.document_parse.source_sha256 == hashlib.sha256(source.body).hexdigest()
    assert result.document_parse.pipeline == "docx"
    assert result.document_parse.table_count == 1
    assert result.document_layout is not None
    assert "table_cells" in result.document_layout.docling_json
    assert result.links[0].url == "https://example.org/appendix.pdf"
    restored = Extracted.model_validate_json(result.model_dump_json())
    assert restored == result
    raw = result.model_dump(by_alias=True)
    raw["text"] += " invented claim"
    with pytest.raises(ValidationError):
        Extracted.model_validate(raw)
    assert not {"docling", "torch", "transformers"} & set(sys.modules)


def test_native_chinese_docx_is_not_translated(tmp_path):
    body = docx(
        "臺灣港務公司公布高雄港碼頭建設與海運基礎設施投資計畫。"
        "研究報告說明港口貨物運輸、航道安全和新碼頭工程，並提供原始資料來源。"
        "地方政府與研究人員檢視相關證據，交通部持續公布工程進度。"
    )
    result = asyncio.run(DocumentExtractor(config(tmp_path)).extract(page(body)))
    assert "高雄港" in result.text and result.language == "zh"


def test_standard_pdf_requires_pinned_offline_artifacts_not_a_native_fallback(tmp_path):
    from tests.test_document_models import artifacts, model_policy

    with pytest.raises(ValidationError):
        config(tmp_path, pdf_pipeline="standard")
    models = model_policy(ocr=None)
    root, entries = artifacts(tmp_path, models)
    cfg = config(
        tmp_path,
        pdf_pipeline="standard",
        artifacts_directory=str(root),
        artifacts=entries,
        pdf_models=models.model_dump(by_alias=True),
    )
    (root / entries[0]["path"]).unlink()
    with pytest.raises(ChimeraRefused) as exc:
        DocumentExtractor(cfg)
    assert exc.value.code == RefusalCode.ADAPTER_CONTRACT


def test_bad_type_or_oversized_source_refuses_before_worker_state(tmp_path):
    client = DocumentExtractor(config(tmp_path, max_input_bytes=10))
    with pytest.raises(ChimeraRefused):
        asyncio.run(client.extract(page()))
    with pytest.raises(ChimeraRefused) as exc:
        asyncio.run(client.extract(page(b"<html></html>", "text/html")))
    assert exc.value.code == RefusalCode.CONTENT_TYPE_UNWANTED
    assert not (tmp_path / "worker").exists()


def test_corrupt_archive_and_expansion_limits_refuse(tmp_path):
    for source, changes in [
        (b"not a zip", {}),
        (docx(), {"max_archive_entries": 1}),
        (docx(), {"max_expanded_bytes": 100}),
    ]:
        with pytest.raises(ChimeraRefused):
            asyncio.run(DocumentExtractor(config(tmp_path, **changes)).extract(page(source)))


def test_mismatched_run_binding_is_refused(tmp_path):
    client = DocumentExtractor(config(tmp_path))
    with pytest.raises(ChimeraRefused):
        client.validate_config(config(tmp_path, max_pages=3))


def test_timeout_releases_worker_slot(tmp_path):
    client = DocumentExtractor(config(tmp_path, timeout_seconds=0.001, max_workers=1))
    with pytest.raises(ChimeraRefused) as exc:
        asyncio.run(client.extract(page()))
    assert exc.value.code == RefusalCode.BUDGET_EXHAUSTED


def native_pdf():
    """A real PDF with a cross-reference table and native text; no mock converter."""
    content = (
        b"BT /F1 12 Tf 50 750 Td (Port infrastructure report) Tj 0 -20 Td "
        b"(The port authority describes maritime transport and terminal construction.) Tj "
        b"0 -20 Td (Analysts reviewed evidence of shipping capacity and investment.) Tj ET"
    )
    objects = (
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        f"<< /Length {len(content)} >>\nstream\n".encode() + content + b"\nendstream",
    )
    data = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, value in enumerate(objects, 1):
        offsets.append(len(data))
        data.extend(f"{number} 0 obj\n".encode() + value + b"\nendobj\n")
    xref = len(data)
    data.extend(b"xref\n0 6\n0000000000 65535 f \n")
    for offset in offsets:
        data.extend(f"{offset:010d} 00000 n \n".encode())
    data.extend(f"trailer\n<< /Size 6 /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    return bytes(data)


def test_real_native_pdf_explicitly_reports_no_layout_model_tables(tmp_path):
    source = page(native_pdf(), "application/pdf")
    result = asyncio.run(DocumentExtractor(config(tmp_path)).extract(source))
    assert "terminal construction" in result.text
    assert result.document_parse.pipeline == "native"
    assert result.document_parse.page_count == 1
    assert result.document_parse.table_count == 0
    assert result.document_parse.artifact_manifest_digest is None
    assert result.language == "en"


def test_document_conversion_is_in_collection_ledger_and_reader(tmp_path):
    from chimera.doubles import FakeJudge, FakeRoute, KeywordScorer
    from chimera.fetch import FetchLadder
    from chimera.loop import GoalLoop
    from chimera.models import Goal, Harvest, Scope

    source = page()

    class DocxRoute(FakeRoute):
        async def attempt(self, request):
            return source

    cfg = config(tmp_path)
    result = asyncio.run(
        GoalLoop(
            config=cfg,
            fetcher=FetchLadder((DocxRoute(),)),
            extractor=DocumentExtractor(cfg),
            scorer=KeywordScorer(),
            judge=FakeJudge(),
        ).run(
            Goal(text="port infrastructure", seeds=(source.url,)),
            Scope(
                allowed_hosts=("example.org",),
                max_depth=0,
                content_types=(source.content_type,),
            ),
        )
    )
    assert len(result.documents) == 1
    assert any(
        row.event == "extraction" and row.document_parse is not None for row in result.ledger
    )
    assert Harvest.model_validate_json(result.model_dump_json()) == result
    raw = result.model_dump(by_alias=True)
    raw["receipt"]["effective_config"]["document_extraction"]["max_pages"] += 1
    with pytest.raises(ValidationError):
        Harvest.model_validate(raw)


def test_cancel_reaps_exact_owned_child_and_next_document_can_run(tmp_path, monkeypatch):
    client = DocumentExtractor(config(tmp_path, max_workers=1))
    created = []
    create = asyncio.create_subprocess_exec

    async def capture(*args, **kwargs):
        process = await create(*args, **kwargs)
        created.append(process)
        return process

    monkeypatch.setattr(asyncio, "create_subprocess_exec", capture)

    async def exercise():
        task = asyncio.create_task(client.extract(page()))
        while not created:
            await asyncio.sleep(0)
        task.cancel()
        with pytest.raises(asyncio.CancelledError):
            await task
        assert created[0].returncode is not None
        assert (await client.extract(page())).document_parse.pipeline == "docx"

    asyncio.run(exercise())
